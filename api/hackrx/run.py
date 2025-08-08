# api/hackrx/run.py
import os, io, json, tempfile, requests, time
from typing import List, Dict
from fastapi import FastAPI, Request, Header, HTTPException
from pypdf import PdfReader
import docx
from email import policy
from email.parser import BytesParser
import numpy as np

# LLM/embedding libs
import openai
from sentence_transformers import SentenceTransformer

# optional vector DBs
try:
    import pinecone
except Exception:
    pinecone = None
try:
    import faiss
except Exception:
    faiss = None

app = FastAPI()

# --- CONFIG from env ---
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
OPENAI_CHAT_MODEL = os.getenv("OPENAI_CHAT_MODEL", "gpt-4")
EMBEDDING_BACKEND = os.getenv("EMBEDDING_BACKEND", "openai")  # "openai" or "sbert"
PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
PINECONE_ENV = os.getenv("PINECONE_ENV")
PINECONE_INDEX = os.getenv("PINECONE_INDEX", "hackrx-index")
TEAM_TOKEN = os.getenv("TEAM_TOKEN", "938fac4d986e2fca94b1c3ae1a2dac27cef31338a6b6ec5033ee7c6a36418a3c")
TOP_K = int(os.getenv("TOP_K", "5"))
SBERT_MODEL = os.getenv("SBERT_MODEL", "all-MiniLM-L6-v2")

if OPENAI_API_KEY:
    openai.api_key = OPENAI_API_KEY

# Optional sbert init
sbert = None
if EMBEDDING_BACKEND == "sbert":
    sbert = SentenceTransformer(SBERT_MODEL)

# Pinecone init if available
use_pinecone = bool(PINECONE_API_KEY and pinecone is not None)
if use_pinecone:
    pinecone.init(api_key=PINECONE_API_KEY, environment=PINECONE_ENV)

# Local FAISS placeholders
FAISS_INDEX = None
FAISS_DOCS = []

# ---------------- HELPERS ----------------
def download_bytes(url: str) -> bytes:
    r = requests.get(url, timeout=20)
    r.raise_for_status()
    return r.content

def parse_pdf_bytes(b: bytes):
    reader = PdfReader(io.BytesIO(b))
    chunks = []
    for pnum, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            chunks.append({"page": pnum, "text": text})
    return chunks

def parse_docx_bytes(b: bytes):
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    tmp.write(b); tmp.flush(); tmp.close()
    doc = docx.Document(tmp.name)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    return [{"page":1, "text": "\n".join(texts)}]

def parse_eml_bytes(b: bytes):
    msg = BytesParser(policy=policy.default).parsebytes(b)
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                body += part.get_content()
    else:
        body = msg.get_content()
    return [{"page":1, "text": body}]

def detect_and_parse(url:str, raw:bytes):
    if url.lower().endswith(".pdf"):
        return parse_pdf_bytes(raw)
    if url.lower().endswith(".docx"):
        return parse_docx_bytes(raw)
    if url.lower().endswith(".eml"):
        return parse_eml_bytes(raw)
    # fallback to pdf parse
    return parse_pdf_bytes(raw)

def chunk_text(text: str, chunk_size=1200, overlap=200):
    text = text.replace("\r", " ")
    chunks = []
    i = 0
    L = len(text)
    while i < L:
        j = min(L, i + chunk_size)
        chunk = text[i:j]
        chunks.append(chunk.strip())
        i = j - overlap
    return [c for c in chunks if c.strip()]

def embed_texts(texts: List[str]) -> List[List[float]]:
    if EMBEDDING_BACKEND == "openai" and OPENAI_API_KEY:
        resp = openai.Embedding.create(model="text-embedding-3-small", input=texts)
        return [r["embedding"] for r in resp["data"]]
    else:
        global sbert
        if sbert is None:
            sbert = SentenceTransformer(SBERT_MODEL)
        embs = sbert.encode(texts, convert_to_numpy=True, show_progress_bar=False)
        norms = np.linalg.norm(embs, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        embs = embs / norms
        return [e.tolist() for e in embs]

def upsert_to_index(docs: List[Dict], emb_list: List[List[float]]):
    global FAISS_INDEX, FAISS_DOCS
    if use_pinecone:
        idx = pinecone.Index(PINECONE_INDEX)
        vectors = []
        for i, emb in enumerate(emb_list):
            meta = {"page": docs[i].get("page"), "source": docs[i].get("source", "")}
            vectors.append((str(i), emb, meta))
        idx.upsert(vectors=vectors)
    else:
        if faiss is None:
            return
        dim = len(emb_list[0])
        xb = np.array(emb_list).astype("float32")
        if FAISS_INDEX is None:
            FAISS_INDEX = faiss.IndexFlatIP(dim)
            FAISS_DOCS = docs.copy()
            FAISS_INDEX.add(xb)
        else:
            FAISS_INDEX.add(xb); FAISS_DOCS.extend(docs)

def query_index(q_emb: List[float], top_k:int=5):
    if use_pinecone:
        idx = pinecone.Index(PINECONE_INDEX)
        res = idx.query(vector=q_emb, top_k=top_k, include_metadata=True)
        out = []
        for m in res["matches"]:
            out.append({"id": m["id"], "score": m["score"], "meta": m.get("metadata", {})})
        return out
    else:
        if FAISS_INDEX is None or FAISS_INDEX.ntotal == 0:
            return []
        xq = np.array([q_emb]).astype("float32")
        D, I = FAISS_INDEX.search(xq, top_k)
        results = []
        for idx, score in zip(I[0], D[0]):
            results.append({"id": str(idx), "score": float(score), "meta": FAISS_DOCS[int(idx)]})
        return results

def build_prompt(question: str, evidence_texts: List[Dict]) -> str:
    ctx = "\n---\n".join([f"[E{i+1}] (page {d.get('page','?')}) {d.get('text')[:1000]}" for i,d in enumerate(evidence_texts)])
    prompt = (
        "You are an expert insurance policy analyst. Use ONLY the evidence provided to answer the question. "
        "If the answer is not found in the evidence, reply 'NOT_IN_DOCUMENT'. Keep the answer concise (max 60 words). "
        "Also return a one-line confidence (high/medium/low) and list evidence tags used.\n\n"
        f"QUESTION: {question}\n\nEVIDENCE:\n{ctx}\n\nFormat strictly as:\n<ANSWER>\n\nEvidence: [E1,E2]\nConfidence: <high|medium|low>\n"
    )
    return prompt

def call_llm(prompt: str) -> str:
    if OPENAI_API_KEY:
        r = openai.ChatCompletion.create(
            model=OPENAI_CHAT_MODEL,
            messages=[{"role":"system","content":"You are a concise, precise assistant."},
                      {"role":"user","content":prompt}],
            temperature=0.0,
            max_tokens=400
        )
        return r["choices"][0]["message"]["content"].strip()
    else:
        return "NOT_IN_DOCUMENT\n\nEvidence: []\nConfidence: low"

# ------------- API ROUTE -------------
@app.post("/hackrx/run")
async def hackrx_run(req: Request, authorization: str = Header(None)):
    start = time.time()
    body = await req.json()
    if authorization is None or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing Bearer token")
    token = authorization.split("Bearer ")[1].strip()
    if token != TEAM_TOKEN:
        raise HTTPException(status_code=403, detail="Invalid token")

    docs_url = body.get("documents")
    questions = body.get("questions", [])
    if not docs_url or not questions:
        raise HTTPException(status_code=400, detail="documents and questions required")

    try:
        raw = download_bytes(docs_url)
        parsed = detect_and_parse(docs_url, raw)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Doc download/parse failed: {str(e)}")

    docs = []
    for p in parsed:
        page_no = p.get("page", 1)
        text = p.get("text","")
        for chunk in chunk_text(text, chunk_size=1200, overlap=200):
            docs.append({"page": page_no, "text": chunk, "source": docs_url})

    texts = [d["text"] for d in docs]
    if len(texts) == 0:
        raise HTTPException(status_code=400, detail="No text extracted from document")
    embeddings = embed_texts(texts)
    upsert_to_index(docs, embeddings)

    answers = []
    for q in questions:
        q_emb = embed_texts([q])[0]
        hits = query_index(q_emb, top_k=TOP_K)
        evidence_texts = []
        for h in hits[:TOP_K]:
            if 'meta' in h and isinstance(h['meta'], dict) and 'text' in h['meta']:
                evidence_texts.append({"page": h['meta'].get("page"), "text": h['meta'].get("text")})
            else:
                try:
                    idx = int(h['id'])
                    evidence_texts.append(FAISS_DOCS[idx])
                except Exception:
                    continue
        prompt = build_prompt(q, evidence_texts)
        resp = call_llm(prompt)
        answers.append(resp.splitlines()[0].strip())
    elapsed = time.time() - start
    return {"answers": answers}
